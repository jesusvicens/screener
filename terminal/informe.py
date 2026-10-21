"""Informe Word (.docx) con todo lo que la terminal tiene de una empresa: dossier, examen, notas, bitácora, estrategia,
KPIs, trimestres y valoración. Es lo que se guarda en la carpeta de OneDrive de cada empresa."""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1B, 0x2A, 0x41)
MUTED = RGBColor(0x6B, 0x72, 0x80)


# ---------------------------------------------------------------- markdown → Word (lo que usa el dossier)
def _runs(par, text: str):
    """Negritas (**x**), cursivas (*x*) y enlaces [t](url) dentro de un párrafo."""
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        else:
            par.add_run(part)


def _table(doc, rows: list[list[str]]):
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c.strip()) or not c.strip() for c in r)]
    if not rows:
        return
    n = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n)
    t.style = "Light Grid Accent 1"
    for i, r in enumerate(rows):
        for j in range(n):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            _runs(p, (r[j] if j < len(r) else "").strip())
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def add_markdown(doc, md: str, base_level: int = 2):
    lines = (md or "").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _table(doc, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            lvl = min(base_level + len(m.group(1)) - 3, 4) if len(m.group(1)) >= 3 else base_level
            doc.add_heading(m.group(2).replace("**", ""), level=max(lvl, 1))
        elif re.match(r"^[-*•]\s+", s):
            indent = (len(ln) - len(ln.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
            _runs(p, re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            indent = (len(ln) - len(ln.lstrip())) // 2
            p = doc.add_paragraph(style="List Number 2" if indent else "List Number")
            _runs(p, re.sub(r"^\d+[.)]\s+", "", s))
        elif s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _runs(p, s.lstrip("> "))
        else:
            p = doc.add_paragraph()
            _runs(p, s)
        i += 1


def _note(doc, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def _toc(doc):
    """Índice de Word (se rellena al abrir: clic derecho → Actualizar campos)."""
    p = doc.add_paragraph()
    run = p.add_run()
    for tag, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        elif txt:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
            run._r.append(el)
        else:
            t = OxmlElement("w:t")
            t.text = "Índice: en Word, clic derecho aquí → «Actualizar campos»."
            run._r.append(t)


def _pct(x):
    return "—" if x is None else f"{x * 100:.1f} %".replace(".", ",")


def _num(x, d=2):
    return "—" if x is None else f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------------------------------------------------------- informe
def build(ticker: str, data: dict, name: str | None = None) -> bytes:
    from . import exam as E
    T = ticker.upper()
    name = name or (data.get("resumen") or {}).get("Empresa") or T
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    h = doc.add_heading(f"{name} ({T})", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _note(doc, f"Análisis de la terminal · actualizado en la app: {data.get('actualizado', '—')} · generado {datetime.now():%d/%m/%Y %H:%M}. "
               "Este documento lo escribe la terminal automáticamente: los cambios que hagas aquí se sobrescribirán; edita en la app.")
    _toc(doc)

    r = data.get("resumen") or {}
    if r:
        doc.add_heading("Resumen de la valoración", level=1)
        rows = [["Dato", "Valor"],
                ["Precio al guardar", f"{_num(r.get('Precio'))} $"], ["PER", f"{_num(r.get('PER'), 1)}x"],
                ["Fase", str(r.get("Fase") or "—")], ["Nota financiera", _num(r.get("Nota financiera"))],
                ["Nota final (con cualitativo)", _num(r.get("Nota final"))], ["Valor estimado", f"{_num(r.get('Valor estimado'))} $"],
                ["Margen de seguridad", _pct(r.get("Margen de seguridad"))], ["Rentabilidad implícita", _pct(r.get("Rentab. implícita"))],
                ["Veredicto", str(r.get("Veredicto") or "—")]]
        _table(doc, rows)
        sup = data.get("supuestos")
        if sup:
            _note(doc, "Supuestos propios de los escenarios: " + "; ".join(f"{k}: " + ", ".join(f"{kk}={vv}" for kk, vv in v.items()) for k, v in sup.items()))
        else:
            _note(doc, "Escenarios: propuesta automática de la app (sin supuestos propios guardados).")

    if data.get("dossier"):
        doc.add_heading("Dossier cualitativo", level=1)
        add_markdown(doc, data["dossier"], base_level=2)

    ans, mot = data.get("examen") or {}, data.get("cualitativo_motivos") or {}
    qs = [q for q in E.QUESTIONS if q[1] in E.QUALITATIVE_BLOCKS and q[0] in ans]
    if qs:
        doc.add_heading("Examen cualitativo", level=1)
        rows = [["Bloque", "Pregunta", "Respuesta", "Motivo"]]
        for q in qs:
            p = ans.get(q[0])
            opts = q[3] if len(q) > 3 else None
            a = "No aplica" if p == 0 else (opts[p - 1] if (opts and p and 1 <= p <= len(opts)) else str(p))
            rows.append([q[1], q[2], f"{a} ({p}/3)" if p else a, mot.get(q[0], "")])
        _table(doc, rows)

    if data.get("notas") or data.get("seguimiento"):
        doc.add_heading("Tesis y seguimiento", level=1)
        if data.get("notas"):
            doc.add_heading("Tesis", level=2)
            add_markdown(doc, data["notas"])
        if data.get("seguimiento"):
            doc.add_heading("Seguimiento", level=2)
            add_markdown(doc, data["seguimiento"])

    if data.get("bitacora"):
        doc.add_heading("Bitácora", level=1)
        for e in sorted(data["bitacora"], key=lambda x: x.get("fecha", ""), reverse=True):
            p = doc.add_paragraph()
            rr = p.add_run(f"{e.get('fecha', '')} · {e.get('tipo', '')}" + (f" · precio {_num(e['precio'])} $" if e.get("precio") else ""))
            rr.bold = True
            add_markdown(doc, e.get("texto", ""))

    g = data.get("guias") or {}
    if g.get("estrategia_anual"):
        doc.add_heading("Estrategia año a año (presentaciones de resultados)", level=1)
        if g.get("estrategia_resumen"):
            add_markdown(doc, g["estrategia_resumen"])
        for y in g["estrategia_anual"][::-1]:
            doc.add_heading(f"{y['año']} · {y.get('titulo', '')}", level=2)
            if y.get("contexto"):
                _note(doc, y["contexto"])
            if y.get("estrategia"):
                p = doc.add_paragraph()
                p.add_run("Estrategia que presentan").bold = True
                for x in y["estrategia"]:
                    _runs(doc.add_paragraph(style="List Bullet"), x)
            if y.get("mensajes"):
                p = doc.add_paragraph()
                p.add_run("Qué dicen").bold = True
                for x in y["mensajes"]:
                    _runs(doc.add_paragraph(style="List Bullet"), x)
            for lab, k in (("Lo que cumplieron", "cumplimiento"), ("Tono", "tono")):
                if y.get(k):
                    p = doc.add_paragraph()
                    p.add_run(f"{lab}: ").bold = True
                    _runs(p, y[k])
    if g.get("kpis_lectura"):
        doc.add_heading("KPIs y segmentos: dicho vs. hecho", level=1)
        add_markdown(doc, g["kpis_lectura"], base_level=2)
    if g.get("trimestres"):
        doc.add_heading("Resultados trimestre a trimestre", level=1)
        for q in g["trimestres"][::-1]:
            doc.add_heading(f"{q['trimestre']} · publicado {q.get('publicado', '')}", level=2)
            if q.get("titular"):
                _runs(doc.add_paragraph(), f"**{q['titular']}**")
            if q.get("resumen"):
                add_markdown(doc, q["resumen"])
            if q.get("tono"):
                p = doc.add_paragraph()
                p.add_run("Tono: ").bold = True
                _runs(p, q["tono"])
            links = [x for x in (q.get("nota_prensa"), q.get("transcripcion_url")) if x]
            if links:
                _note(doc, " · ".join(links))
    if data.get("comparables"):
        doc.add_heading("Comparables", level=1)
        doc.add_paragraph(", ".join(data["comparables"]))
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def filename(ticker: str, name: str | None = None) -> str:
    base = re.sub(r"[\\/:*?\"<>|]", "", (name or ticker)).strip()
    base = re.sub(r"\s*,?\s*(INC\.?|CORP\.?|CORPORATION|CO\.?)$", "", base, flags=re.I).strip() or ticker
    return f"{base.title() if base.isupper() else base} ({ticker.upper()}) - Análisis terminal.docx"
