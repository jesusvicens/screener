"""Genera tu Excel de análisis (misma plantilla que «ADOBE - analisis financer 2026.xlsx») para cualquier empresa.

Se parte de templates/plantilla_analisis.xlsx (tu Excel, con todas sus fórmulas) y se rellenan:
- PyG, Balance y Cash Flow (últimos 10 ejercicios, el último en la columna M).
- Precios medios por ejercicio y precio actual (Modelo EPS), parámetros y fase (Escenarios), tus supuestos
  escritos a mano, la parte de riesgo (probabilidades y precios de prueba, caída intermedia, perfiles) y el
  test de exposición.
- Examen: bloque cualitativo con tus respuestas y el examen financiero nuevo (4 bloques, 21 preguntas) con
  sus fórmulas, desplegables y el dato que justifica cada propuesta.
Excel recalcula todas las fórmulas al abrir el archivo.
"""
from __future__ import annotations

import io
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import exam as E
from .riesgo import DEFAULT_DROP, DEFAULT_MIN_GUARANTEE, DEFAULT_PROFILES, DEFAULT_TEST_PRICES, DEFAULT_TEST_PROBS

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "plantilla_analisis.xlsx"
COLS = [get_column_letter(c) for c in range(4, 14)]  # D..M

PYG = {8: "revenue", 9: "cogs", 10: "gross_profit", 11: "rnd", 12: "sga", 13: None, 14: "opex", 15: "operating_income", 16: "nonop",
       17: "pretax", 18: "taxes", 19: "net_income", 20: None, 21: "net_income", 22: None, 23: "net_income", 24: "ebitda",
       25: "operating_income", 26: "shares_basic", 27: "shares_diluted", 28: "eps_basic", 29: "eps_diluted"}
BAL = {5: "cash_sti", 6: "receivables", 7: "inventory", 8: "prepaid", 9: "other_current_assets", 10: "current_assets", 11: "ppe",
       12: "lt_investments", 13: "goodwill_intangibles", 14: "other_lt_assets", 15: "lt_assets", 16: "total_assets",
       17: "current_liabilities", 18: "lt_debt", 19: "other_lt_liabilities", 20: "lt_liabilities", 21: "total_liabilities", 22: None,
       23: "retained_earnings", 24: "aoci", 25: None, 26: "equity", 27: "liabilities_equity"}
CF = {6: "cf_net_income", 7: "da", 8: None, 9: None, 10: "chg_receivables", 11: "chg_inventory", 12: "chg_payables", 13: None, 14: None,
      15: "ocf", 16: "capex", 17: None, 18: "acquisitions", 19: None, 20: None, 21: None, 22: None, 23: "icf", 24: "debt_net", 25: None,
      26: "debt_net", 27: "equity_net", 28: "equity_net", 29: "dividends", 30: None, 31: "fcf_fin", 32: "net_cash_flow", 33: "sbc",
      34: "dividends"}
SC_ROWS = {"g1": 72, "g10": 73, "gm": 74, "om": 75, "buyback": 76, "per_exit": 77, "prob": 78}
SC_COLS = {"Optimista": "C", "Base": "D", "Adverso": "E", "Desastre": "F"}


def _fill_statement(ws, mapping, header_row, f, years):
    yrs = years[-10:]
    offset = 10 - len(yrs)
    for i, col in enumerate(COLS):
        y = yrs[i - offset] if i >= offset else None
        ws[f"{col}{header_row}"].value = y
        for row, key in mapping.items():
            v = f.get(key, y) if (key and y is not None) else None
            ws[f"{col}{row}"].value = None if v is None else round(float(v), 4)
    return yrs, offset


def _replace_names(wb, name, ticker):
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=50):
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    v = c.value.replace("ADOBE (ADBE)", f"{name} ({ticker})").replace("ADOBE", name).replace("ADBE", ticker)
                    if v != c.value:
                        c.value = v


def _style_from(dst, src):
    dst.font, dst.fill, dst.border, dst.alignment, dst.number_format = (copy(src.font), copy(src.fill), copy(src.border),
                                                                        copy(src.alignment), src.number_format)


def build_workbook(ca, work: dict | None = None) -> bytes:
    """ca: pipeline.CompanyAnalysis; work: lo que hay en pantalla (supuestos, examen, riesgo, notas)."""
    work = work or {}
    f, p = ca.fin, ca.params
    name = f.name.upper().replace(" INC.", "").replace(" INC", "").replace(",", "").strip()
    wb = openpyxl.load_workbook(TEMPLATE)
    _replace_names(wb, name, f.ticker)

    # --- Estados financieros
    yrs, offset = _fill_statement(wb["PyG"], PYG, 7, f, f.years)
    _fill_statement(wb["Balance"], BAL, 4, f, f.years)
    _fill_statement(wb["Cash Flow"], CF, 5, f, f.years)
    es = wb["Escenarios"]
    for i, col in enumerate(COLS):
        es[f"{col}16"].value = str(yrs[i - offset]) if i >= offset else ""
    wb["ROIC"]["A1"].value = f"{name} · ROIC y capital invertido (histórico {yrs[0]}-{yrs[-1]})"
    wb["ROIC"]["B30"].value = int(p.rd_life)

    # --- Modelo EPS
    me = wb["Modelo EPS"]
    for i in range(10):
        y = yrs[i - offset] if i >= offset else None
        me[f"C{2 + i}"].value = round(ca.avg_prices[y], 2) if (y in ca.avg_prices) else None
    me["B21"].value = yrs[-1] + 10
    me["D25"].value = round(p.price, 2) if p.price else None
    me["C25"].value = p.r

    # --- Escenarios: parámetros, fase y supuestos escritos a mano
    es["B6"].value = p.price_date or ""
    es["B7"].value, es["B8"].value, es["B9"].value, es["B10"].value, es["B11"].value = p.r, int(p.n), p.mos, p.g_lt, int(p.years_to_margin)
    if work.get("fase"):
        es["B56"].value = work["fase"]
    for scn, vals in (work.get("supuestos") or {}).items():
        for fld, v in (vals or {}).items():
            if fld in SC_ROWS and scn in SC_COLS and v is not None:
                es[f"{SC_COLS[scn]}{SC_ROWS[fld]}"].value = float(v)
    es["B164"].value, es["B165"].value, es["B166"].value = p.patrimonio, p.importe, p.deuda_posicion or 0

    # --- Riesgo (Modelo EPS filas 43-126)
    rk = work.get("riesgo") or {}
    tp = {**DEFAULT_TEST_PROBS, **(rk.get("test_probs") or {})}
    for i, scn in enumerate(["Optimista", "Base", "Adverso", "Desastre"]):
        me[f"J{55 + i}"].value = tp[scn]
    tpx = (rk.get("test_prices") or DEFAULT_TEST_PRICES) + [None, None, None]
    for i, r in enumerate((99, 100, 101)):
        me[f"B{r}"].value = tpx[i]
    me["C104"].value = rk.get("drop", DEFAULT_DROP)
    me["C105"].value = rk.get("min_guarantee", DEFAULT_MIN_GUARANTEE)
    profs = (rk.get("profiles") or DEFAULT_PROFILES)[:5]
    for i in range(5):
        r = 108 + i
        pr = profs[i] if i < len(profs) else None
        me[f"B{r}"].value = pr.get("Inversor") if pr else None
        me[f"C{r}"].value = float(pr.get("Peso") or 0) if pr else None
        me[f"D{r}"].value = float(pr.get("Deuda / capital") or 0) if pr else None
        me[f"L{r}"].value = pr.get("Lectura") if pr else None

    if offset:
        _tolerate_missing_years(wb)
    # Rentabilidad anual con valor futuro negativo: -100 % en lugar de #NUM! (igual que en la terminal)
    for ref in [f"E{r}" for r in range(55, 60)] + [f"G{r}" for r in range(33, 40)]:
        v = me[ref].value
        if isinstance(v, str) and v.startswith("=") and not v.upper().startswith("=IFERROR("):
            me[ref].value = f"=IFERROR({v[1:]},-1)"
    _write_exam(wb, ca, work)
    wb.calculation.fullCalcOnLoad = True
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _tolerate_missing_years(wb):
    """Con menos de 10 ejercicios las primeras columnas quedan vacías: las fórmulas históricas que las usan
    devuelven vacío en lugar de #DIV/0! (las medias, mínimos y máximos ignoran esos huecos)."""
    def wrap(ws, rows, cols=None):
        for row in ws.iter_rows(min_row=rows[0], max_row=rows[1]):
            for c in row:
                if cols and c.column_letter not in cols:
                    continue
                v = c.value
                if isinstance(v, str) and v.startswith("=") and not v.upper().startswith("=IFERROR("):
                    c.value = f'=IFERROR({v[1:]},"")'
    wrap(wb["Modelo EPS"], (2, 11))
    wrap(wb["Escenarios"], (17, 25))
    wrap(wb["Escenarios"], (27, 52), cols={"B"})
    wrap(wb["ROIC"], (4, 41))


def _write_exam(wb, ca, work):
    ex, li = wb["Examen"], wb["Listas"]
    answers = dict(ca.answers or {})
    answers.update(work.get("examen") or {})
    q_style = {c: ex[f"{c}22"] for c in "CDEFGH"}
    head_style = ex["C22"]
    # Cualitativo: filas 2-21 (mismas preguntas y fórmulas de tu plantilla)
    for qid, block, text, opts in E.QUESTIONS:
        if block in E.QUALITATIVE_BLOCKS:
            r = int(qid[1:])
            pts = answers.get(qid)
            ex[f"G{r}"].value = opts[pts - 1] if pts else None
    # Limpia el bloque financiero antiguo (y sus celdas combinadas)
    for mr in list(ex.merged_cells.ranges):
        if mr.max_row >= 22:
            ex.unmerge_cells(str(mr))
    for r in range(22, 80):
        for col in "ABCDEFGHIJ":
            ex[f"{col}{r}"].value = None
    for r in range(22, 60):
        for col in "ABCDE":
            li[f"{col}{r}"].value = None
    li["E1"].value = "Opción «No aplica»"
    # Financiero nuevo: 4 bloques, 21 preguntas
    r = 22
    block_cells = []
    dv_lists = []
    for block in E.FINANCIAL_BLOCKS:
        qs = [q for q in E.QUESTIONS if q[1] == block]
        first = r
        ex[f"C{r}"].value = block
        ex[f"C{r}"].value = f"{block}\n({E.BLOCK_SUBTITLE.get(block, '')})"
        _style_from(ex[f"C{r}"], head_style)
        ex[f"E{r}"].value = f'=IFERROR(SUM(H{first}:H{first + len(qs) - 1})/COUNT(H{first}:H{first + len(qs) - 1})/3*10,"n.a.")'
        block_cells.append(f"E{r}")
        for qid, _b, text, opts in qs:
            pts = answers.get(qid)
            ex[f"F{r}"].value = text
            ex[f"G{r}"].value = opts[pts - 1] if pts else (E.NA if pts == 0 else None)
            o1, o2, o3 = (o.replace('"', '""') for o in opts)
            ex[f"H{r}"].value = (f'=IF(G{r}="{E.NA}","",IF(G{r}="{o1}",1,IF(G{r}="{o2}",2,IF(G{r}="{o3}",3,0))))')
            sug = ca.suggestions.get(qid)
            ex[f"I{r}"].value = sug[1] if sug else None
            ex[f"I{r}"].font = Font(size=8, color="666666")
            ex[f"I{r}"].alignment = Alignment(wrap_text=False)
            for c in "FGH":
                _style_from(ex[f"{c}{r}"], q_style[c])
            # Listas para el desplegable
            li[f"A{r}"].value, li[f"B{r}"].value, li[f"C{r}"].value, li[f"D{r}"].value, li[f"E{r}"].value = text, *opts, E.NA
            dv_lists.append((f"G{r}", f"Listas!$B${r}:$E${r}"))
            r += 1
        _style_from(ex[f"E{first}"], q_style["E"])
        ex.merge_cells(f"C{first}:D{r - 1}")
        ex.merge_cells(f"E{first}:E{r - 1}")
        ex[f"C{first}"].alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ex[f"E{first}"].alignment = Alignment(vertical="center", horizontal="center")
    last = r - 1
    ex.merge_cells(f"A2:A{last}")
    ex.merge_cells(f"B22:B{last}")
    ex["B22"].alignment = Alignment(vertical="center", horizontal="center")
    ex["B22"].value = "=AVERAGE(" + ",".join(block_cells) + ")"
    ex["I1"].value = "Dato que justifica la propuesta (terminal)"
    ex["I1"].font = Font(bold=True, size=9)
    ex.column_dimensions["I"].width = 70
    # Notas
    notes = "\n\n".join(x for x in [work.get("notas") or ca.saved.get("notas") or "", work.get("seguimiento") or ""] if x)
    ex["K1"].value = "Notas:"
    ex["K2"].value = notes or None
    ex["K2"].alignment = Alignment(wrap_text=True, vertical="top")
    # Desplegables (se re-crean: los de la plantilla original no se conservan al editarla)
    ex.data_validations.dataValidation = []
    for qid, block, _t, _o in E.QUESTIONS:
        if block in E.QUALITATIVE_BLOCKS:
            rr = int(qid[1:])
            dv_lists.append((f"G{rr}", f"Listas!$B${rr}:$D${rr}"))
    for cell, rng in dv_lists:
        dv = DataValidation(type="list", formula1=rng, allow_blank=True)
        dv.add(cell)
        ex.add_data_validation(dv)
    es = wb["Escenarios"]
    dv = DataValidation(type="list", formula1="Listas!$H$2:$H$7", allow_blank=False)
    dv.add("B56")
    es.add_data_validation(dv)
    highlight = PatternFill("solid", fgColor="FFF2CC")
    ex["A2"].fill = highlight
