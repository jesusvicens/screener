"""Importa a la terminal las respuestas del examen y tus notas de uno de tus Excel de análisis.

Uso:  python importar_excel.py "ADOBE - analisis financer 2026.xlsx" ADBE
Lee la hoja «Examen» (columna G = respuesta), la nota de seguimiento (columna K), tus parámetros de
«Escenarios» y la parte de riesgo de «Modelo EPS» (probabilidades de prueba, precios de prueba,
caída intermedia, garantía mínima y perfiles de inversor), y guarda
data/analisis/<TICKER>.json (o en GitHub si tienes configurados los secrets).
"""
import sys

import openpyxl

from terminal import exam, storage


def importar(path: str, ticker: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Examen"]
    data = storage.load(ticker) or {}
    answers = dict(data.get("examen") or {})
    for qid, _block, _text, opts in exam.QUESTIONS:
        v = ws[f"G{int(qid[1:])}"].value
        if v in opts:
            answers[qid] = opts.index(v) + 1
    notes = [str(ws[f"K{r}"].value) for r in range(1, ws.max_row + 1) if ws[f"K{r}"].value]
    data["examen"] = answers
    if notes:
        data["seguimiento"] = "\n".join(notes)
    try:
        es = wb["Escenarios"]
        params = dict(data.get("params") or {})
        for key, cell in [("r", "B7"), ("n", "B8"), ("mos", "B9"), ("g_lt", "B10"), ("years_to_margin", "B11")]:
            if isinstance(es[cell].value, (int, float)):
                params[key] = es[cell].value
        data["params"] = params
    except KeyError:
        pass
    try:
        me = wb["Modelo EPS"]
        rk = dict(data.get("riesgo") or {})
        names = ["Optimista", "Base", "Adverso", "Desastre"]
        tp = {n: me[f"J{55 + i}"].value for i, n in enumerate(names)}
        if all(isinstance(v, (int, float)) for v in tp.values()):
            rk["test_probs"] = tp
        if isinstance(me["C104"].value, (int, float)):
            rk["drop"] = me["C104"].value
        if isinstance(me["C105"].value, (int, float)):
            rk["min_guarantee"] = me["C105"].value
        tpx = [me[f"B{r}"].value for r in (99, 100, 101) if isinstance(me[f"B{r}"].value, (int, float))]
        if tpx:
            rk["test_prices"] = tpx
        profs = []
        for r in range(108, 113):
            if me[f"B{r}"].value:
                profs.append({"Inversor": me[f"B{r}"].value, "Peso": me[f"C{r}"].value or 0, "Deuda / capital": me[f"D{r}"].value or 0,
                              "Lectura": me[f"L{r}"].value or ""})
        if profs:
            rk["profiles"] = profs
        data["riesgo"] = rk
    except KeyError:
        pass
    print(storage.save(ticker, data))
    return data


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    importar(sys.argv[1], sys.argv[2])
